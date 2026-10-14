"""Extract ERP workbook headers and their current QMS destination fields."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Dict, Iterable, List, Tuple

import xlrd
from openpyxl import load_workbook

from qms_erp_sync import (
    BUSINESS_SOURCE_TABLE_BY_FILE,
    DATE_FIELDS,
    DECIMAL_FIELDS,
    FIELD_MAP,
    INTEGER_FIELDS,
    WORK_ORDER_DB_FIELDS,
    excel_header_identity,
    load_erp_field_registry,
    normalize_header,
    is_work_order_sheet,
    source_table_name,
    unique_headers,
)


ROOT = Path(__file__).resolve().parent
SOURCE_DIRECTORY = ROOT / "data" / "BI数据源"
OUTPUT = ROOT / "outputs" / "excel_database_field_mapping_data.json"
EXCEL_SUFFIXES = (".xls", ".xlsx")

def excel_column(index: int) -> str:
    result = ""
    value = index
    while value:
        value, remainder = divmod(value - 1, 26)
        result = chr(ord("A") + remainder) + result
    return result


def iter_workbook_headers(path: Path) -> Iterable[Tuple[str, List[Any]]]:
    if path.name.lower().endswith(".xlsx"):
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                yield sheet.title, list(next(sheet.iter_rows(values_only=True), ()))
        finally:
            workbook.close()
        return

    workbook = xlrd.open_workbook(str(path), on_demand=True)
    try:
        for sheet_name in workbook.sheet_names():
            sheet = workbook.sheet_by_name(sheet_name)
            yield sheet_name, sheet.row_values(0) if sheet.nrows else []
    finally:
        workbook.release_resources()


def table_hint(file_name: str) -> str:
    for key, target in BUSINESS_SOURCE_TABLE_BY_FILE:
        if key in file_name:
            return target
    return "qms_erp_unclassified"


def direct_targets(header: str) -> List[Tuple[str, str, str]]:
    normalized = normalize_header(header)
    api_to_database = {api_field: database_field for database_field, api_field in WORK_ORDER_DB_FIELDS}
    results: List[Tuple[str, str, str]] = []
    for api_field, source_header in FIELD_MAP.items():
        if normalize_header(source_header) != normalized:
            continue
        if api_field == "workOrderNo":
            results.append(("qms_work_order", "work_order_no", "生产指令号作为工单唯一编号"))
        elif api_field == "batchNo":
            results.append(("qms_production_batch", "batch_no", "管制卡号作为生产批次编号"))
        else:
            database_field = api_to_database.get(api_field)
            if database_field:
                if api_field in DATE_FIELDS:
                    rule = "Excel 日期转换为 YYYY-MM-DD"
                elif api_field in INTEGER_FIELDS:
                    rule = "转换为整数"
                elif api_field in DECIMAL_FIELDS:
                    rule = "转换为十进制数值"
                else:
                    rule = "按文本/原值写入"
                results.append(("qms_work_order", database_field, rule))

    special = {
        normalize_header("入库/已结批"): ("qms_work_order", "warehouse_status", "按入库、结批等状态转换"),
        normalize_header("加急订单"): ("qms_work_order", "priority", "加急值转换为 URGENT，否则 NORMAL"),
        normalize_header("本卡工作板数"): ("qms_work_order", "total_panel_quantity", "优先作为总工作板数"),
        normalize_header("C板总PNL"): ("qms_work_order", "total_panel_quantity", "缺少本卡工作板数时参与总板数计算"),
    }
    if normalized in special:
        results.append(special[normalized])
    if normalized in {
        normalize_header("材质"), normalize_header("板厚"), normalize_header("铜箔"),
    }:
        results.append(("qms_work_order", "specification", "与材质、板厚、铜箔拼接生成规格"))
    if normalized == normalize_header("投料数量"):
        results.append(("qms_production_batch", "planned_quantity", "生产批次计划数量"))
    return results


def main() -> None:
    registry = load_erp_field_registry()
    mapping_rows: List[Dict[str, str]] = []
    source_summary: List[Dict[str, Any]] = []
    files = sorted(
        path for path in SOURCE_DIRECTORY.rglob("*")
        if path.is_file() and path.name.lower().endswith(EXCEL_SUFFIXES)
    )
    for path in files:
        relative_path = path.relative_to(SOURCE_DIRECTORY).as_posix()
        hint = table_hint(path.name)
        sheet_count = 0
        column_count = 0
        for sheet_name, raw_headers in iter_workbook_headers(path):
            sheet_count += 1
            work_order_sheet = "生产指示资料" in path.name and is_work_order_sheet(raw_headers)
            fixed_table = source_table_name(relative_path, sheet_name, work_order_sheet)
            unique = unique_headers(raw_headers)
            for index, header in enumerate(unique, start=1):
                source_header = excel_header_identity(raw_headers[index - 1])
                direct = direct_targets(header) if "生产指示资料" in path.name else []
                field = registry.resolve(
                    "qms_work_order" if work_order_sheet else fixed_table,
                    relative_path, sheet_name, source_header, header, index,
                )
                erp_field = field.database_column
                if work_order_sheet:
                    # All production-instruction columns are real qms_work_order
                    # columns.  Existing QMS mappings remain alongside them.
                    targets = [("qms_work_order", erp_field, "原 Excel 值直接写入实体字段")] + direct
                    seen = set()
                    unique_targets = []
                    for target in targets:
                        identity = (target[0], target[1])
                        if identity not in seen:
                            seen.add(identity)
                            unique_targets.append(target)
                    current_business_table = "\n".join(target[0] for target in unique_targets)
                    current_business_field = "\n".join(target[1] for target in unique_targets)
                    rule = "\n".join(target[2] for target in unique_targets)
                else:
                    current_business_table = fixed_table or "qms_erp_unclassified"
                    current_business_field = erp_field
                    rule = "原 Excel 值直接写入实体字段"
                mapping_rows.append({
                    "Excel 文件": relative_path,
                    "工作表": sheet_name,
                    "Excel 列": excel_column(index),
                    "Excel 原始列名": source_header or "（空列名）",
                    "Excel 唯一列名": header,
                    "英文数据库字段": field.database_column,
                    "数据库字段注释": field.column_comment,
                    "旧数据库字段": field.legacy_database_column,
                    "当前写入表": current_business_table,
                    "当前写入字段": current_business_field,
                    "当前业务表": current_business_table,
                    "当前业务字段": current_business_field,
                    "转换规则": rule,
                    "说明": (
                        "生产指示资料通过 QMS 工单和生产批次映射写入"
                        if work_order_sheet else "固定业务表；来源文件和工作表保存在 source_file、sheet_name"
                    ),
                })
                column_count += 1
        source_summary.append({
            "Excel 文件": relative_path,
            "工作表数": sheet_count,
            "列数": column_count,
            "固定写入表": hint if "生产指示资料" not in path.name else "qms_work_order / qms_production_batch",
        })

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps({
        "mapping_rows": mapping_rows,
        "source_summary": source_summary,
        "source_table_columns": [
            ["tenant_code", "租户编码"],
            ["source_file", "ERP 相对文件路径"],
            ["sheet_name", "工作表名称"],
            ["source_row_no", "Excel 行号"],
            ["business_key", "来源业务识别键"],
            ["source_key", "来源记录哈希键"],
            ["row_hash", "Excel 行数据指纹"],
            ["erp_<english_name>", "每个 Excel 列对应一个注册表英文实体字段；原始中文列名保留在字段注释中"],
            ["status", "ACTIVE / INACTIVE"],
        ],
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"提取完成：{len(files)} 个文件，{len(mapping_rows)} 个 Excel 列。")


if __name__ == "__main__":
    main()
