"""Behavioral coverage for the authoritative ERP field registry."""

import json
import tempfile
import unittest
import importlib
from dataclasses import FrozenInstanceError, replace
from pathlib import Path
from decimal import Decimal
from unittest.mock import patch, Mock

from openpyxl import Workbook

import qms_erp_sync as sync


REGISTRY_PATH = Path(__file__).resolve().parents[1] / "erp_field_registry.json"


class ErpFieldRegistryTests(unittest.TestCase):
    def test_registry_matches_nested_daily_path_and_normalizes_repeated_excel_suffix(self):
        field = sync.ErpFieldDefinition(
            "qms_purchase_order", "采购订单资料.xls.xls", "Test", "采购单号", "采购单号",
            "erp_purchase_order_no", "erp_采购单号", "Excel列名：采购单号", 1,
        )

    def test_registry_accepts_daily_export_collapsing_a_leading_backslash(self):
        field = sync.ErpFieldDefinition(
            "qms_customer_part", "客户料号资料.xls.xls", "Test", "\\\\库存面积", "\\\\库存面积",
            "erp_inventory_area", "erp_\\\\库存面积", "Excel列名：\\\\库存面积", 24,
        )
        registry = sync.ErpFieldRegistry([field])
        self.assertEqual(
            field,
            registry.resolve(
                "qms_customer_part", "2026-9-17/客户料号资料.xls", "Test",
                "\\库存面积", "\\库存面积", 24,
            ),
        )
        registry = sync.ErpFieldRegistry([field])
        self.assertEqual(
            field,
            registry.resolve(
                "qms_purchase_order", "2026-9-17/采购订单资料.XLS", "Test",
                "采购单号", "采购单号", 1,
            ),
        )

    def test_header_identity_preserves_nonblank_strings_and_cell_text_semantics(self):
        for header in ("\r加工方式", "加工方式\r", "\x85加工方式", "加工方式\x85",
                       "\x80加工方式", "加工方式\x7f", "\u2003加工方式\u00a0", "null"):
            with self.subTest(header=header):
                self.assertEqual(sync.excel_header_identity(header), header)
        self.assertEqual(sync.as_text("\r加工方式\x85"), "加工方式")
        self.assertEqual(sync.as_text("null"), "")
        for value, expected in [(None, ""), (True, "是"), (False, "否"), (1.0, "1"), (Decimal("1.20"), "1.2")]:
            self.assertEqual(sync.excel_header_identity(value), expected)

    def test_header_identity_keeps_distinct_raw_keys_and_blank_placeholders(self):
        headers = ["\r加工方式", "加工方式\r", "\x85加工方式", "加工方式", "加工方式",
                   None, "", " \t\r\n\x85\u2003\u00a0 "]
        self.assertEqual(sync.unique_headers(headers), [
            "\r加工方式", "加工方式\r", "\x85加工方式", "加工方式", "加工方式（重复列2）",
            "列6", "列7", "列8",
        ])
        for header in (None, "", " \t\r\n\x85\u2003\u00a0 "):
            self.assertEqual(sync.excel_header_identity(header), "")
        for header in headers[:4]:
            self.assertEqual(sync.legacy_database_column_name(header), "erp_加工方式")

    def test_header_identity_flows_through_generator_xlsx_scan_and_mapping(self):
        import generate_field_mapping_data as mapping
        generator = importlib.import_module("generate_erp_field_registry")
        headers = ["\x85加工方式", "加工方式\n", "\u2003加工方式\u00a0", " \t\x85 ", "\x80加工方式"]
        expected_sources = headers[:3] + [""] + headers[4:]
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook = Workbook()
            workbook.active.title = "Test"
            workbook.active.append(headers)
            workbook.active.append(["first", "second", "third", "blank", "fifth"])
            workbook.save(root / "客户订单资料.xlsx")
            workbook.close()
            data = generator.build_registry(root)
            self.assertEqual([f["excel_source_header"] for f in data["fields"]], expected_sources)
            self.assertEqual([f["excel_unique_header"] for f in data["fields"]], headers[:3] + ["列4"] + headers[4:])
            self.assertEqual([f["column_comment"] for f in data["fields"]], [
                "Excel列名： 加工方式", "Excel列名：加工方式 ", "Excel列名：\u2003加工方式\u00a0",
                "Excel列名：", "Excel列名： 加工方式",
            ])
            self.assertEqual([f["legacy_database_column"] for f in data["fields"]], [
                "erp_加工方式", "erp_加工方式", "erp_加工方式", "erp_列4", "erp_\x80加工方式",
            ])
            registry = sync.ErpFieldRegistry(sync.ErpFieldDefinition(**row) for row in data["fields"])
            with patch.object(sync, "load_erp_field_registry", return_value=registry):
                pairs = list(sync.iter_all_source_records({"source_directory": str(root)}, lambda _: None, sync.SourceScanStats()))
            self.assertEqual(list(pairs[0][0].field_values.values()), ["first", "second", "third", "blank", "fifth"])
            self.assertEqual([f.excel_source_header for f in pairs[0][0].field_definitions.values()], expected_sources)
            output = root / "mapping.json"
            with patch.object(mapping, "SOURCE_DIRECTORY", root), patch.object(mapping, "OUTPUT", output), \
                    patch.object(mapping, "load_erp_field_registry", return_value=registry):
                mapping.main()
            rows = json.loads(output.read_text(encoding="utf-8"))["mapping_rows"]
            self.assertEqual([row["Excel 原始列名"] for row in rows], headers[:3] + ["（空列名）"] + headers[4:])
            self.assertEqual([row["数据库字段注释"] for row in rows], [f["column_comment"] for f in data["fields"]])

    def test_header_identity_xls_scan_preserves_leading_and_trailing_cr(self):
        headers = ["\r加工方式", "加工方式\r", "\x85加工方式"]
        fields = [sync.ErpFieldDefinition(
            "qms_customer_order", "客户订单资料.xls", "Test", header, header,
            "erp_processing_method_" + str(index), "erp_加工方式", sync.erp_column_comment(header), index,
        ) for index, header in enumerate(headers, 1)]
        registry = sync.ErpFieldRegistry(fields)
        sheet = Mock(nrows=2)
        sheet.row_values.side_effect = lambda index: headers if index == 0 else ["a", "b", "c"]
        workbook = Mock(datemode=0)
        workbook.sheet_names.return_value = ["Test"]
        workbook.sheet_by_name.return_value = sheet
        root = Path("D:/header-identity-test")
        with patch.object(sync, "candidate_files", return_value=[root / "客户订单资料.xls"]), \
                patch.object(sync.xlrd, "open_workbook", return_value=workbook), \
                patch.object(sync, "load_erp_field_registry", return_value=registry):
            pairs = list(sync.iter_all_source_records({"source_directory": str(root)}, lambda _: None, sync.SourceScanStats()))
        raw, _ = pairs[0]
        self.assertEqual(list(raw.field_values.values()), ["a", "b", "c"])
        self.assertEqual([f.excel_source_header for f in raw.field_definitions.values()], headers)

    def test_xls_date_cells_are_written_as_iso_dates_without_converting_plain_numbers(self):
        headers = ["采购日期", "数量"]
        fields = [
            sync.ErpFieldDefinition(
                "qms_purchase_order", "采购订单资料.xls.xls", "Test", header, header,
                column, sync.legacy_database_column_name(header), sync.erp_column_comment(header), index,
            )
            for index, (header, column) in enumerate([
                ("采购日期", "erp_purchase_date"),
                ("数量", "erp_quantity"),
            ], start=1)
        ]
        registry = sync.ErpFieldRegistry(fields)
        sheet = Mock(nrows=2)
        sheet.row_values.side_effect = lambda index: headers if index == 0 else [42195.0, 42195.0]
        sheet.cell_type.side_effect = lambda row, column: (
            sync.xlrd.XL_CELL_DATE if (row, column) == (1, 0) else sync.xlrd.XL_CELL_NUMBER
        )
        workbook = Mock(datemode=0)
        workbook.sheet_names.return_value = ["Test"]
        workbook.sheet_by_name.return_value = sheet
        root = Path("D:/xls-date-test")
        with patch.object(sync, "candidate_files", return_value=[root / "采购订单资料.xls.xls"]), \
                patch.object(sync.xlrd, "open_workbook", return_value=workbook), \
                patch.object(sync, "load_erp_field_registry", return_value=registry):
            pairs = list(sync.iter_all_source_records({"source_directory": str(root)}, lambda _: None, sync.SourceScanStats()))
        raw, _ = pairs[0]
        self.assertEqual(raw.field_values, {
            "erp_purchase_date": "2015-07-10",
            "erp_quantity": "42195",
        })

    def test_xls_date_conversion_preserves_time_and_workbook_date_system(self):
        self.assertEqual(
            sync.xls_safe_cell(42195.5, sync.xlrd.XL_CELL_DATE, 0),
            "2015-07-10 12:00:00",
        )
        self.assertEqual(
            sync.xls_safe_cell(1.0, sync.xlrd.XL_CELL_DATE, 1),
            "1904-01-02",
        )
        self.assertEqual(
            sync.xls_safe_cell(42195.0, sync.xlrd.XL_CELL_NUMBER, 0),
            42195,
        )
        self.assertIsNone(sync.xls_safe_cell(2.0, sync.xlrd.XL_CELL_DATE, 0))

    def test_xlsx_dates_use_the_same_database_friendly_format(self):
        self.assertEqual(sync.xlsx_safe_cell(sync.datetime(2015, 7, 10)), "2015-07-10")
        self.assertEqual(
            sync.xlsx_safe_cell(sync.datetime(2015, 7, 10, 12, 34, 56)),
            "2015-07-10 12:34:56",
        )
        self.assertEqual(sync.xlsx_safe_cell(42195.0), 42195)

    def test_comment_control_characters_are_collapsed_without_changing_normal_headers(self):
        for value in list(range(0x20)) + list(range(0x7f, 0xa0)):
            with self.subTest(codepoint=value):
                self.assertEqual(sync.erp_column_comment("加工" + chr(value) + "方式"), "Excel列名：加工 方式")
        self.assertEqual(sync.erp_column_comment("加工\r\n\t\x00\x85方式"), "Excel列名：加工 方式")
        self.assertEqual(sync.erp_column_comment("采购单号"), "Excel列名：采购单号")
        self.assertEqual(sync.erp_column_comment("入库面积  (料号)"), "Excel列名：入库面积  (料号)")

    def test_comment_control_sanitation_preserves_raw_scope_and_legacy_name(self):
        source = "加工\r\n方式"
        field = sync.ErpFieldDefinition(
            "qms_customer_order", "客户订单资料.xls", "Test", source, source,
            "erp_processing_method", "erp_加工方式", "Excel列名：加工 方式", 1,
        )
        registry = sync.ErpFieldRegistry([field])
        self.assertEqual(registry.resolve(field.table_name, field.source_file, "Test", source, source, 1), field)
        self.assertEqual(field.excel_source_header, "加工\r\n方式")
        self.assertEqual(field.excel_unique_header, "加工\r\n方式")
        self.assertEqual(field.legacy_database_column, "erp_加工方式")
        with self.assertRaisesRegex(sync.SyncError, "注释"):
            sync.ErpFieldRegistry([replace(field, column_comment="Excel列名：加工\r\n方式")])

    def test_comment_control_sanitation_flows_from_real_workbook_to_mapping(self):
        import generate_field_mapping_data as mapping
        generator = importlib.import_module("generate_erp_field_registry")
        source = "订单数量\n<生产通知单>"
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook = Workbook()
            workbook.active.title = "Test"
            workbook.active.append([source])
            workbook.save(root / "客户订单资料.xlsx")
            workbook.close()
            data = generator.build_registry(root)
            field = data["fields"][0]
            self.assertEqual(field["column_comment"], "Excel列名：订单数量 <生产通知单>")
            self.assertEqual(field["excel_source_header"], source)
            self.assertEqual(field["excel_unique_header"], source)
            self.assertEqual(field["legacy_database_column"], "erp_订单数量<生产通知单>")
            registry = sync.ErpFieldRegistry([sync.ErpFieldDefinition(**field)])
            output = root / "mapping.json"
            with patch.object(mapping, "SOURCE_DIRECTORY", root), patch.object(mapping, "OUTPUT", output), \
                    patch.object(mapping, "load_erp_field_registry", return_value=registry):
                mapping.main()
            row = json.loads(output.read_text(encoding="utf-8"))["mapping_rows"][0]
            self.assertEqual(row["数据库字段注释"], "Excel列名：订单数量 <生产通知单>")
            self.assertEqual(row["Excel 原始列名"], source)

    def test_generator_uses_semantics_placeholders_duplicates_and_rejects_unknown_terms(self):
        generator = importlib.import_module("generate_erp_field_registry")
        for source, unique, index, expected in [
            ("采购单号", "采购单号", 1, "erp_purchase_order_no"),
            ("税率%", "税率%", 9, "erp_tax_rate_percent"),
            ("其他17", "其他17", 83, "erp_other_17"),
            ("其他bi1", "其他bi1", 29, "erp_bi_1"),
            ("其它mc3", "其它mc3", 9, "erp_mc_3"),
            ("", "列9", 9, "erp_blank_column_9"),
            ("采购单号", "采购单号（重复列2）", 2, "erp_purchase_order_no_duplicate_2"),
        ]:
            with self.subTest(header=unique):
                self.assertEqual(generator.english_column_name(source, unique, index), expected)
        with self.assertRaisesRegex(sync.SyncError, "词典"):
            generator.english_column_name("未知龘", "未知龘", 1)

    def test_registry_resolves_a_readable_english_column_and_chinese_comment(self):
        registry = sync.load_erp_field_registry(REGISTRY_PATH)
        field = registry.resolve(
            "qms_purchase_order", "采购订单资料.xls.xls", "Test",
            "采购单号", "采购单号", 1,
        )
        self.assertEqual(field.database_column, "erp_purchase_order_no")
        self.assertEqual(field.column_comment, "Excel列名：采购单号")
        self.assertEqual(field.legacy_database_column, "erp_采购单号")
        with self.assertRaises(FrozenInstanceError):
            field.database_column = "erp_changed"

    def test_unregistered_header_raises_actionable_registry_error(self):
        registry = sync.load_erp_field_registry(REGISTRY_PATH)
        with self.assertRaisesRegex(sync.SyncError, "字段注册表.*未登记的新字段"):
            registry.resolve(
                "qms_purchase_order", "采购订单资料.xls.xls", "Test",
                "未登记的新字段", "未登记的新字段", 999,
            )

    def test_source_scope_and_header_metadata_cannot_resolve_a_different_field(self):
        registry = sync.load_erp_field_registry(REGISTRY_PATH)
        for file_name, sheet, source, unique, index in [
            ("another.xls", "Test", "采购单号", "采购单号", 1),
            ("采购订单资料.xls.xls", "Other", "采购单号", "采购单号", 1),
            ("采购订单资料.xls.xls", "Test", "采购编号", "采购单号", 1),
            ("采购订单资料.xls.xls", "Test", "采购单号", "采购单号", 2),
        ]:
            with self.subTest(file=file_name, sheet=sheet, source=source, index=index):
                with self.assertRaisesRegex(sync.SyncError, "字段注册表"):
                    registry.resolve("qms_purchase_order", file_name, sheet, source, unique, index)

    def test_loader_rejects_unsafe_names_tables_collisions_and_inconsistent_comments(self):
        field = {
            "table_name": "qms_purchase_order", "source_file": "采购订单资料.xls.xls",
            "sheet_name": "Test", "excel_source_header": "采购单号",
            "excel_unique_header": "采购单号", "column_index": 1,
            "database_column": "erp_purchase_order_no", "legacy_database_column": "erp_采购单号",
            "column_comment": "Excel列名：采购单号",
        }
        cases = [
            [dict(field, table_name="users")],
            [dict(field, database_column="erp_采购单号")],
            [dict(field, database_column="erp_" + "a" * 61)],
            [dict(field, column_comment="incorrect")],
            [dict(field, legacy_database_column="erp_wrong")],
            [dict(field, column_index=0)],
            [field, dict(field, database_column="erp_another")],
            [field, dict(field, source_file="other.xls")],
        ]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "registry.json"
            for fields in cases:
                with self.subTest(fields=fields):
                    path.write_text(json.dumps({"schema_version": 1, "fields": fields}), encoding="utf-8")
                    with self.assertRaisesRegex(sync.SyncError, "字段注册表"):
                        sync.load_erp_field_registry(path)

    def test_missing_or_malformed_registry_is_an_actionable_sync_error(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "missing.json"
            with self.assertRaisesRegex(sync.SyncError, "字段注册表"):
                sync.load_erp_field_registry(path)
            path.write_text("{", encoding="utf-8")
            with self.assertRaisesRegex(sync.SyncError, "字段注册表"):
                sync.load_erp_field_registry(path)

    def test_default_registry_loader_uses_pyinstaller_bundle_resource(self):
        """The frozen EXE must read the JSON bundled under PyInstaller's resource root."""
        with tempfile.TemporaryDirectory() as directory:
            bundle = Path(directory)
            (bundle / "erp_field_registry.json").write_text(
                REGISTRY_PATH.read_text(encoding="utf-8"), encoding="utf-8")
            with patch.object(sync, "__file__", str(bundle / "app" / "qms_erp_sync.py")), \
                    patch.object(sync.sys, "frozen", True, create=True), \
                    patch.object(sync.sys, "_MEIPASS", str(bundle), create=True):
                registry = sync.load_erp_field_registry()
            self.assertGreater(len(registry.fields), 1)

    def test_generator_reads_real_workbooks_and_assigns_stable_source_scoped_names(self):
        generator = importlib.import_module("generate_erp_field_registry")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for file_name in ("原料出库单资料.xlsx", "补料资料.xlsx"):
                workbook = Workbook()
                workbook.active.title = "Test"
                workbook.active.append(["数量", "数量", None, "其他17"])
                workbook.save(root / file_name)
                workbook.close()
            data = generator.build_registry(root)
            self.assertEqual(len(data["fields"]), 8)
            self.assertEqual(len(data["collision_resolutions"]), 8)
            fields = [sync.ErpFieldDefinition(**row) for row in data["fields"]]
            registry = sync.ErpFieldRegistry(fields)
            first = registry.resolve("qms_material_issue", "原料出库单资料.xlsx", "Test", "数量", "数量", 1)
            second = registry.resolve("qms_material_issue", "补料资料.xlsx", "Test", "数量", "数量", 1)
            self.assertNotEqual(first.database_column, second.database_column)
            base_fields = [sync.ErpFieldDefinition(
                field.table_name, field.source_file, field.sheet_name, field.excel_source_header,
                field.excel_unique_header,
                generator.english_column_name(field.excel_source_header, field.excel_unique_header, field.column_index),
                field.legacy_database_column, field.column_comment, field.column_index,
            ) for field in fields]
            reversed_result = generator.disambiguate_fields(list(reversed(base_fields)))
            self.assertEqual(fields, list(reversed(reversed_result)))
            self.assertTrue(all(len(field.database_column.encode("utf-8")) <= 64 for field in fields))

    def test_mapping_export_uses_registry_names_and_chinese_metadata(self):
        import generate_field_mapping_data as mapping
        generator = importlib.import_module("generate_erp_field_registry")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook = Workbook()
            workbook.active.title = "Test"
            workbook.active.append(["采购单号", "税率%"])
            workbook.save(root / "采购订单资料.xlsx")
            workbook.close()
            data = generator.build_registry(root)
            registry = sync.ErpFieldRegistry(sync.ErpFieldDefinition(**row) for row in data["fields"])
            output = root / "mapping.json"
            with patch.object(mapping, "SOURCE_DIRECTORY", root), patch.object(mapping, "OUTPUT", output), \
                    patch.object(mapping, "load_erp_field_registry", return_value=registry):
                mapping.main()
            rows = json.loads(output.read_text(encoding="utf-8"))["mapping_rows"]
            self.assertEqual(rows[0]["英文数据库字段"], "erp_purchase_order_no")
            self.assertEqual(rows[0]["数据库字段注释"], "Excel列名：采购单号")
            self.assertEqual(rows[0]["旧数据库字段"], "erp_采购单号")
            self.assertEqual(rows[1]["当前写入字段"], "erp_tax_rate_percent")

    def test_real_mapping_generator_exports_english_identifier_and_chinese_comment(self):
        """A delivery mapping must expose the registry's BI-ready field metadata."""
        import generate_field_mapping_data as mapping

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "real-mapping.json"
            with patch.object(mapping, "OUTPUT", output):
                mapping.main()

            rows = json.loads(output.read_text(encoding="utf-8"))["mapping_rows"]
            purchase_order_no = next(
                row for row in rows
                if row["Excel 文件"] == "采购订单资料.xls.xls"
                and row["工作表"] == "Test"
                and row["Excel 原始列名"] == "采购单号"
            )
            self.assertRegex(purchase_order_no["英文数据库字段"], r"^erp_[a-z][a-z0-9_]*$")
            self.assertEqual(purchase_order_no["数据库字段注释"], "Excel列名：采购单号")
            self.assertTrue(purchase_order_no["旧数据库字段"].startswith("erp_"))
            self.assertNotEqual(
                purchase_order_no["英文数据库字段"], purchase_order_no["旧数据库字段"])


if __name__ == "__main__":
    unittest.main()
