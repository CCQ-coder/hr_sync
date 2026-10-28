"""Recovery and overwritten-file behavior for the ERP directory synchronizer."""

import os
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pymysql
from openpyxl import Workbook

import qms_erp_sync as sync


class FileCheckpointTests(unittest.TestCase):
    def test_upgrade_replays_latest_unchanged_complete_exports_once(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            baseline = root / "baseline"
            daily = root / "daily"
            baseline.mkdir()
            daily.mkdir()
            customer = baseline / "客户料号资料.xls"
            material_old = baseline / "原料资料明细.xls"
            material_new = daily / "2026-10-23" / "原料资料明细.xls"
            receipt = daily / "2026-10-23" / "成品入库单资料.xls"
            for path in (customer, material_old, material_new, receipt):
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_bytes(path.as_posix().encode())
            config = {
                "initialization_directory": str(baseline),
                "incremental_directory": str(daily),
            }
            _, checkpoints = sync.build_source_scan_plan(config, {})
            previous = {**checkpoints, sync.SOURCE_IDENTITY_VERSION_STATE_KEY: "5"}
            groups, _ = sync.build_source_scan_plan(config, previous)
            self.assertEqual(
                [("initialization", [customer]), ("incremental", [material_new, receipt])],
                [(group.role, list(group.files)) for group in groups],
            )
            previous["complete-snapshot-version"] = "1"
            groups, _ = sync.build_source_scan_plan(config, previous)
            self.assertEqual([], groups)

    def test_changed_older_initialization_snapshot_replays_latest_same_table(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            older = root / "2026-10-22" / "原料资料明细.xls"
            latest = root / "2026-10-23" / "原料资料明细.xls"
            for path in (older, latest):
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_bytes(b"initial")
            config = {"initialization_directory": str(root)}
            _, checkpoints = sync.build_source_scan_plan(config, {})
            previous = {
                **checkpoints,
                sync.SOURCE_IDENTITY_VERSION_STATE_KEY: "5",
                "complete-snapshot-version": "1",
            }
            older.write_bytes(b"changed")
            groups, _ = sync.build_source_scan_plan(config, previous)
            self.assertEqual([older, latest], list(groups[0].files))

    def test_complete_snapshot_hard_deletes_only_missing_current_tenant_erp_rows(self) -> None:
        class Cursor:
            def __init__(self):
                self.rows = {
                    ("customer-b", "ERP5:keep"), ("customer-b", "ERP5:missing"),
                    ("customer-b", "manual-entry"), ("customer-c", "ERP5:missing"),
                }
                self.rowcount = 0
                self.selected = []

            def execute(self, statement, parameters):
                if statement.lstrip().startswith("SELECT"):
                    self.assert_select(statement, parameters)
                    self.selected = [(key,) for tenant, key in self.rows
                                     if tenant == parameters[0] and key.startswith("ERP5:")]
                else:
                    self.assert_delete(statement, parameters)
                    before = len(self.rows)
                    self.rows.difference_update({(parameters[0], key) for key in parameters[1:]})
                    self.rowcount = before - len(self.rows)

            def assert_select(self, statement, parameters):
                assert "`tenant_code` = %s" in statement
                assert "`source_key` LIKE 'ERP5:%%'" in statement
                assert "FORCE INDEX (`uq_427354425f7dfe1e_source`)" in statement
                assert parameters == ("customer-b",)

            def assert_delete(self, statement, parameters):
                assert statement.lstrip().startswith("DELETE FROM `qms_material`")
                assert "`tenant_code` = %s" in statement
                assert "`source_key` LIKE 'ERP5:%%'" in statement

            def fetchall(self):
                return self.selected

        client = object.__new__(sync.QmsDatabaseClient)
        client.tenant = "customer-b"
        cursor = Cursor()
        count, deleted_keys = client._delete_missing_snapshot_rows(cursor, "qms_material", {"ERP5:keep"})
        self.assertEqual(1, count)
        self.assertEqual({"ERP5:missing"}, deleted_keys)
        self.assertEqual({
            ("customer-b", "ERP5:keep"), ("customer-b", "manual-entry"),
            ("customer-c", "ERP5:missing"),
        }, cursor.rows)

    def test_complete_snapshot_sql_survives_pymysql_percent_formatting(self) -> None:
        class Cursor:
            def __init__(self):
                self.formatter = pymysql.cursors.Cursor(
                    SimpleNamespace(literal=lambda value: repr(value)),
                )
                self.statements = []
                self.rowcount = 1

            def execute(self, statement, parameters):
                self.statements.append(self.formatter.mogrify(statement, parameters))

            def fetchall(self):
                return [("ERP5:keep",), ("ERP5:missing",)]

        client = object.__new__(sync.QmsDatabaseClient)
        client.tenant = "customer-b"
        cursor = Cursor()
        try:
            deleted, keys = client._delete_missing_snapshot_rows(
                cursor, "qms_material", {"ERP5:keep"},
            )
        except TypeError as exc:
            self.fail(f"PyMySQL could not bind the snapshot query: {exc}")

        self.assertEqual((1, {"ERP5:missing"}), (deleted, keys))
        self.assertEqual(2, len(cursor.statements))
        self.assertTrue(all("`source_key` LIKE 'ERP5:%'" in sql for sql in cursor.statements))
        self.assertTrue(all("`tenant_code` = 'customer-b'" in sql for sql in cursor.statements))

    def test_same_named_exports_in_different_daily_roots_have_distinct_identity(self) -> None:
        class Registry:
            def canonical_headers_for(self, _source_file, headers):
                return headers

            def resolve_sheet(self, *_args):
                return []

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            files = []
            for day, order in (("10-19", "A"), ("10-20", "B")):
                folder = root / day
                folder.mkdir()
                path = folder / "客户订单资料.xlsx"
                workbook = Workbook()
                workbook.active.append(["客户订单号"])
                workbook.active.append([order])
                workbook.save(path)
                files.append((folder, path))

            config = {"source_file_keyword": "生产指示资料"}
            with patch.object(sync, "load_erp_field_registry", return_value=Registry()):
                rows = [
                    next(iter(sync.iter_all_source_records(
                        config, lambda _: None, sync.SourceScanStats(),
                        source_directory=folder, selected_files=(path,),
                    )))[0]
                    for folder, path in files
                ]
            self.assertNotEqual(rows[0].source_key, rows[1].source_key)
            self.assertNotEqual(rows[0].source_file, rows[1].source_file)

            first_config = {"skip_initialization": True, "incremental_directory": str(files[0][0])}
            second_config = {"skip_initialization": True, "incremental_directory": str(files[1][0])}
            _, previous = sync.build_source_scan_plan(first_config, {})
            groups, _ = sync.build_source_scan_plan(second_config, previous)
            self.assertEqual([files[1][1]], list(groups[0].files))

    def test_same_path_overwrite_changes_content_key_and_hash(self) -> None:
        class Registry:
            def canonical_headers_for(self, _source_file, headers):
                return headers

            def resolve_sheet(self, *_args):
                return []

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "客户订单资料.xlsx"

            def write_order(order):
                workbook = Workbook()
                workbook.active.append(["客户订单号"])
                workbook.active.append([order])
                workbook.save(path)

            def scan():
                with patch.object(sync, "load_erp_field_registry", return_value=Registry()):
                    return next(iter(sync.iter_all_source_records(
                        {}, lambda _: None, sync.SourceScanStats(),
                        source_directory=root, selected_files=(path,),
                    )))[0]

            write_order("A")
            before = scan()
            write_order("B")
            after = scan()
            self.assertNotEqual(before.source_key, after.source_key)
            self.assertNotEqual(before.row_hash, after.row_hash)

    def test_same_record_in_two_exports_uses_one_source_key(self) -> None:
        class Registry:
            def canonical_headers_for(self, _source_file, headers):
                return headers

            def resolve_sheet(self, *_args):
                return []

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            paths = []
            for day in ("10-19", "10-20"):
                folder = root / day
                folder.mkdir()
                path = folder / "客户订单资料.xlsx"
                workbook = Workbook()
                workbook.active.append(["客户订单号"])
                workbook.active.append(["A"])
                workbook.save(path)
                paths.append(path)
            with patch.object(sync, "load_erp_field_registry", return_value=Registry()):
                keys = [
                    next(iter(sync.iter_all_source_records(
                        {}, lambda _: None, sync.SourceScanStats(),
                        source_directory=root, selected_files=(path,),
                    )))[0].source_key
                    for path in paths
                ]
            self.assertEqual(keys[0], keys[1])

    def test_same_size_and_mtime_overwrite_is_detected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            workbook = root / "2026-10-20" / "采购订单资料.xls"
            workbook.parent.mkdir()
            workbook.write_bytes(b"first")
            original_stat = workbook.stat()
            config = {"skip_initialization": True, "incremental_directory": str(root)}
            _, previous = sync.build_source_scan_plan(config, {})
            original_fingerprint = sync.directory_fingerprint(root)

            workbook.write_bytes(b"other")
            os.utime(workbook, ns=(original_stat.st_atime_ns, original_stat.st_mtime_ns))

            groups, _ = sync.build_source_scan_plan(config, previous)
            self.assertEqual([workbook], list(groups[0].files))
            self.assertNotEqual(original_fingerprint, sync.directory_fingerprint(root))

    def test_interrupted_baseline_replay_keeps_incremental_files_pending(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            baseline = root / "baseline"
            incremental = root / "incremental"
            baseline.mkdir()
            incremental.mkdir()
            (baseline / "base.xls").write_bytes(b"base")
            (incremental / "delta.xls").write_bytes(b"delta")
            config = {
                "initialization_directory": str(baseline),
                "incremental_directory": str(incremental),
            }
            _, previous = sync.build_source_scan_plan(config, {})
            previous[sync.BASELINE_REPLAY_STATE_KEY] = "1"
            groups, _ = sync.build_source_scan_plan(config, previous)
            self.assertEqual(["incremental"], [group.role for group in groups])

    def test_removed_rows_are_scoped_to_the_overwritten_file(self) -> None:
        first_key = "source-file:incremental:2026-10-20/one.xls"
        second_key = "source-file:incremental:2026-10-20/two.xls"
        state = {
            sync.file_membership_state_key(first_key): json.dumps({
                "raw": {"qms_purchase_order": ["old", "shared"]},
                "work_orders": ["ERP-EXCEL:old", "ERP-EXCEL:shared"],
                "batches": ["old-batch", "shared-batch"],
            }),
            sync.file_membership_state_key(second_key): json.dumps({
                "raw": {"qms_purchase_order": ["shared"]},
                "work_orders": ["ERP-EXCEL:shared"],
                "batches": ["shared-batch"],
            }),
        }
        current = {"raw": {"qms_purchase_order": []}, "work_orders": [], "batches": []}
        self.assertEqual({
            "raw": {"qms_purchase_order": ["old"]},
            "work_orders": ["ERP-EXCEL:old"],
            "batches": ["old-batch"],
        }, sync.removed_file_membership(state, first_key, current))

    def test_deactivation_sql_limits_changes_to_owned_keys_and_erp_orders(self) -> None:
        class Cursor:
            rowcount = 1

            def __init__(self):
                self.calls = []

            def execute(self, statement, parameters):
                self.calls.append((statement, parameters))

        client = object.__new__(sync.QmsDatabaseClient)
        client.tenant = "customer-b"
        cursor = Cursor()
        count = client._deactivate_removed_file_records(cursor, {
            "raw": {"qms_purchase_order": ["old"]},
            "work_orders": ["ERP-EXCEL:old"],
            "batches": ["old-batch"],
        })
        self.assertEqual(3, count)
        self.assertEqual(3, len(cursor.calls))
        self.assertIn("`tenant_code` = %s", cursor.calls[0][0])
        self.assertIn("'SUPERSEDED'", cursor.calls[0][0])
        self.assertEqual(("customer-b", "old"), cursor.calls[0][1])
        self.assertIn("`source_system` = 'ERP_EXCEL'", cursor.calls[1][0])
        self.assertIn("work_order.`source_system` = 'ERP_EXCEL'", cursor.calls[2][0])

    def test_source_table_creation_is_not_repeated_for_every_batch(self) -> None:
        class Cursor:
            def __init__(self):
                self.statements = []

            def execute(self, statement, _parameters=None):
                self.statements.append(statement)

        client = object.__new__(sync.QmsDatabaseClient)
        client._legacy_json_checked = set()
        client._source_table_ready = set()
        cursor = Cursor()
        with patch.object(client, "_read_table_columns", return_value={}), \
                patch.object(client, "_ensure_erp_columns"):
            client._ensure_source_table_schema(cursor, "qms_material", {})
            client._ensure_source_table_schema(cursor, "qms_material", {})
        create_statements = [item for item in cursor.statements if "CREATE TABLE" in item]
        self.assertEqual(1, len(create_statements))

    def test_same_legacy_row_is_reused_without_duplicating_the_import(self) -> None:
        raw = sync.RawSourceRecord(
            source_file="D:/daily/客户订单资料.XLS", sheet_name="Test",
            source_row_no=2, business_key="A#行2", source_key="new-key",
            row_hash="same-hash", field_values={},
            source_table="qms_customer_order", field_definitions={},
        )
        old_key = sync.legacy_source_key("客户订单资料.XLS", "Test", 2)
        existing = {old_key: ("客户订单资料.XLS", "same-hash")}

        self.assertEqual(
            old_key,
            sync.compatible_source_key(raw, "客户订单资料.XLS", existing),
        )

    def test_changed_legacy_row_is_preserved_when_new_day_uses_same_row(self) -> None:
        raw = sync.RawSourceRecord(
            source_file="D:/daily/客户订单资料.XLS", sheet_name="Test",
            source_row_no=2, business_key="B#行2", source_key="new-key",
            row_hash="new-hash", field_values={},
            source_table="qms_customer_order", field_definitions={},
        )
        old_key = sync.legacy_source_key("客户订单资料.XLS", "Test", 2)
        existing = {old_key: ("客户订单资料.XLS", "old-hash")}

        self.assertEqual("new-key", sync.compatible_source_key(raw, "客户订单资料.XLS", existing))

    def test_adopted_legacy_row_updates_in_place_after_same_path_overwrite(self) -> None:
        raw = sync.RawSourceRecord(
            source_file="D:/daily/客户订单资料.XLS", sheet_name="Test",
            source_row_no=2, business_key="B#行2", source_key="new-key",
            row_hash="new-hash", field_values={},
            source_table="qms_customer_order", field_definitions={},
        )
        old_key = sync.legacy_source_key("客户订单资料.XLS", "Test", 2)
        self.assertEqual(old_key, sync.compatible_source_key(
            raw, "客户订单资料.XLS", {old_key: (raw.source_file, "old-hash")},
        ))
        self.assertEqual("new-key", sync.compatible_source_key(
            raw, "客户订单资料.XLS", {old_key: ("D:/other/客户订单资料.XLS", "new-hash")},
        ))

    def test_existing_new_key_takes_precedence_over_matching_legacy_row(self) -> None:
        raw = sync.RawSourceRecord(
            source_file="D:/daily/客户订单资料.XLS", sheet_name="Test",
            source_row_no=2, business_key="B#行2", source_key="new-key",
            row_hash="new-hash", field_values={},
            source_table="qms_customer_order", field_definitions={},
        )
        old_key = sync.legacy_source_key("客户订单资料.XLS", "Test", 2)
        existing = {old_key: ("客户订单资料.XLS", "new-hash"), "new-key": (raw.source_file, "new-hash")}
        self.assertEqual("new-key", sync.compatible_source_key(raw, "客户订单资料.XLS", existing))

    def test_legacy_lookup_is_limited_to_tenant_and_current_file(self) -> None:
        class Cursor:
            def __init__(self):
                self.calls = []

            def execute(self, statement, parameters):
                self.calls.append((statement, parameters))

            def fetchall(self):
                return [("old-key", "客户订单资料.XLS", "same-hash")]

        client = object.__new__(sync.QmsDatabaseClient)
        client.tenant = "customer-b"
        cursor = Cursor()
        result = client._source_records_for_file(
            cursor, "qms_customer_order", "客户订单资料.XLS", "D:/daily/客户订单资料.XLS",
        )
        self.assertEqual({"old-key": ("客户订单资料.XLS", "same-hash")}, result)
        self.assertIn("tenant_code = %s", cursor.calls[0][0])
        self.assertIn("source_file IN (%s, %s)", cursor.calls[0][0])
        self.assertEqual(("customer-b", "客户订单资料.XLS", "D:/daily/客户订单资料.XLS"), cursor.calls[0][1])


class ResumeTests(unittest.TestCase):
    def test_complete_export_deletes_missing_rows_before_commit_and_marks_upgrade(self) -> None:
        events = []

        class Connection:
            def cursor(self):
                return self

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            def commit(self):
                events.append("commit")

            def close(self):
                pass

        class Client:
            def _connect(self):
                return Connection()

            def _upsert_source_records(self, _cursor, _table, records):
                events.append(("upsert", [record.source_key for record in records]))

            def _delete_missing_snapshot_rows(self, _cursor, table, keys):
                events.append(("delete", table, set(keys)))
                return 2, {"ERP5:old"}

            def _deactivate_removed_file_records(self, _cursor, _removed):
                return 0

        def records(_config, _callback, stats, *, source_directory, selected_files):
            stats.files_scanned += 1
            stats.raw_records += 1
            key = selected_files[0].read_text()
            yield sync.RawSourceRecord(
                source_file=selected_files[0].as_posix(), sheet_name="Test", source_row_no=2,
                business_key=key, source_key=f"ERP5:{key}", row_hash=key, field_values={},
                source_table="qms_material", field_definitions={},
            ), None

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            daily = root / "daily"
            daily.mkdir()
            for day, value in (("2026-10-22", "old"), ("2026-10-23", "new")):
                folder = daily / day
                folder.mkdir()
                (folder / "原料资料明细.xls").write_text(value)
            config = {"skip_initialization": True, "incremental_directory": str(daily)}
            with patch.object(sync, "state_path", return_value=root / "state.json"), \
                    patch.object(sync, "iter_all_source_records", side_effect=records), \
                    patch.object(sync, "QmsDatabaseClient", return_value=Client()):
                sync.save_state({sync.SOURCE_IDENTITY_VERSION_STATE_KEY: "5"})
                result = sync.synchronize_database_directory(config, False, lambda _: None)
                state = sync.load_state()
            self.assertEqual([
                "commit",
                ("upsert", ["ERP5:new"]),
                ("delete", "qms_material", {"ERP5:new"}),
                "commit",
            ], events)
            self.assertEqual(2, result.deleted)
            self.assertEqual("1", state["complete-snapshot-version"])
            self.assertNotIn("raw:ERP5:old", state)
            self.assertEqual("new", state["raw:ERP5:new"])

    def test_empty_complete_incremental_export_skips_and_continues(self) -> None:
        events = []
        messages = []

        class Connection:
            def cursor(self):
                return self

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            def commit(self):
                events.append("commit")

            def rollback(self):
                events.append("rollback")

            def close(self):
                pass

        class Client:
            def _connect(self):
                return Connection()

            def _upsert_source_records(self, _cursor, _table, rows):
                events.append(("upsert", rows[0].source_key))

            def _delete_missing_snapshot_rows(self, *_args):
                events.append("delete")
                return 0, set()

            def _deactivate_removed_file_records(self, *_args):
                events.append("deactivate")
                return 0

        def records(_config, _callback, stats, *, source_directory, selected_files):
            stats.files_scanned += 1
            path = selected_files[0]
            value = path.read_text()
            if value == "empty":
                return
            stats.raw_records += 1
            yield sync.RawSourceRecord(
                source_file=path.as_posix(), sheet_name="Test", source_row_no=2,
                business_key=value, source_key=f"ERP5:{value}", row_hash=value,
                field_values={}, source_table=sync.source_table_name(path.name, "", False),
                field_definitions={},
            ), None

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            daily = root / "daily"
            daily.mkdir()
            older = daily / "2026-10-26" / "原料资料明细.xls"
            empty = daily / "2026-10-27" / "原料资料明细.xls"
            later = daily / "2026-10-28" / "采购订单资料.xls"
            for path in (older, empty, later):
                path.parent.mkdir()
                path.write_text("old" if path == older else "empty" if path == empty else "next")
            config = {"skip_initialization": True, "incremental_directory": str(daily)}
            with patch.object(sync, "state_path", return_value=root / "state.json"), \
                    patch.object(sync, "iter_all_source_records", side_effect=records), \
                    patch.object(sync, "QmsDatabaseClient", return_value=Client()):
                sync.save_state({sync.SOURCE_IDENTITY_VERSION_STATE_KEY: "5"})
                result = sync.synchronize_database_directory(config, False, messages.append)
                state = sync.load_state()
                self.assertIn(sync.source_file_state_key("incremental", daily, empty), state)
                self.assertEqual("1", state["complete-snapshot-version"])
                self.assertEqual(0, sync.synchronize_database_directory(config, False, messages.append).source_files)
            self.assertEqual(2, result.source_files)
            self.assertNotIn(("upsert", "ERP5:old"), events)
            self.assertIn(("upsert", "ERP5:next"), events)
            self.assertNotIn("delete", events)
            self.assertNotIn("rollback", events)
            self.assertEqual(1, events.count("deactivate"))
            self.assertTrue(any("跳过" in message and "原料资料明细.xls" in message for message in messages))

    def test_empty_complete_reinitialize_skips_without_deletion(self) -> None:
        events = []

        class Connection:
            def cursor(self):
                return self

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            def commit(self):
                events.append("commit")

            def execute(self, statement):
                events.append(statement)

            def rollback(self):
                events.append("rollback")

            def close(self):
                pass

        class Client:
            def _connect(self):
                return Connection()

            def _clear_target_tables(self, _cursor, _callback):
                events.append("clear")

            def _delete_missing_snapshot_rows(self, *_args):
                raise AssertionError("empty export must not delete rows")

        def no_records(_config, _callback, stats, *, source_directory, selected_files,
                       headers_only=False, on_sheet=None, read_path=None):
            stats.files_scanned += 1
            return iter(())

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "baseline"
            source.mkdir()
            (source / "原料资料明细.xls").write_bytes(b"empty")
            config = {"initialization_directory": str(source)}
            with patch.object(sync, "state_path", return_value=root / "state.json"), \
                    patch.object(sync, "iter_all_source_records", side_effect=no_records), \
                    patch.object(sync, "QmsDatabaseClient", return_value=Client()):
                result = sync.synchronize_database_directory(
                    config, False, lambda _: None, initialize=True,
                )
                self.assertEqual(0, result.deleted)
                self.assertTrue((root / "state.json").exists())
        self.assertEqual([
            "commit", "SET TRANSACTION ISOLATION LEVEL REPEATABLE READ", "clear", "commit",
        ], events)

    def test_failed_latest_snapshot_is_replayed_after_older_file_checkpoint(self) -> None:
        events = []

        class Connection:
            def cursor(self):
                return self

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            def commit(self):
                events.append("commit")

            def rollback(self):
                events.append("rollback")

            def close(self):
                pass

        class Client:
            fail_latest = True

            def _connect(self):
                return Connection()

            def _upsert_source_records(self, _cursor, _table, rows):
                if rows[0].source_key == "ERP5:new" and self.fail_latest:
                    raise sync.SyncError("newer failed")
                events.append(("upsert", rows[0].source_key))

            def _delete_missing_snapshot_rows(self, _cursor, _table, keys):
                events.append(("delete", set(keys)))
                return 0, set()

            def _deactivate_removed_file_records(self, *_args):
                return 0

        def records(_config, _callback, stats, *, source_directory, selected_files):
            stats.files_scanned += 1
            stats.raw_records += 1
            value = selected_files[0].read_text()
            yield sync.RawSourceRecord(
                source_file=selected_files[0].as_posix(), sheet_name="Test", source_row_no=2,
                business_key=value, source_key=f"ERP5:{value}", row_hash=value,
                field_values={}, source_table="qms_material", field_definitions={},
            ), None

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            older = root / "2026-10-22" / "原料资料明细.xls"
            latest = root / "2026-10-23" / "原料资料明细.xls"
            for path, value in ((older, "old"), (latest, "new")):
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_text(value)
            config = {"initialization_directory": str(root)}
            _, checkpoints = sync.build_source_scan_plan(config, {})
            client = Client()
            with patch.object(sync, "state_path", return_value=root / "state.json"), \
                    patch.object(sync, "iter_all_source_records", side_effect=records), \
                    patch.object(sync, "QmsDatabaseClient", return_value=client):
                sync.save_state({
                    **checkpoints, sync.SOURCE_IDENTITY_VERSION_STATE_KEY: "5",
                    "complete-snapshot-version": "1",
                })
                older.write_text("changed")
                with self.assertRaisesRegex(sync.SyncError, "newer failed"):
                    sync.synchronize_database_directory(config, False, lambda _: None)
                self.assertEqual(["commit", "rollback"], events)
                state = sync.load_state()
                self.assertEqual("pending", state["complete-snapshot-version"])
                groups, _ = sync.build_source_scan_plan(config, state)
                self.assertEqual([latest], list(groups[0].files))
                client.fail_latest = False
                sync.synchronize_database_directory(config, False, lambda _: None)
                self.assertEqual("1", sync.load_state()["complete-snapshot-version"])
            self.assertEqual(("delete", {"ERP5:new"}), events[-2])

    def test_file_sync_retains_equal_customer_order_details_without_reconciliation(self) -> None:
        class Connection:
            def cursor(self):
                return self

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            def commit(self):
                if len(client.written_keys) != 2:
                    raise AssertionError("both customer-order details must be written before commit")

            def close(self):
                pass

        class Client:
            def __init__(self):
                self.written_keys = []

            def _connect(self):
                return Connection()

            def _existing_source_tables(self, _cursor):
                return {"qms_customer_order"}

            def _source_records_for_file(self, _cursor, _table, relative, _absolute):
                raise AssertionError("location-based legacy key must not be reused")

            def _upsert_source_records(self, _cursor, _table, records):
                self.written_keys.extend(record.source_key for record in records)

            def _deactivate_removed_file_records(self, _cursor, _removed):
                return 0

        def fake_records(_config, _callback, stats, *, source_directory, selected_files):
            stats.files_scanned += 1
            stats.raw_records += 2
            for row_no in (2, 3):
                yield sync.RawSourceRecord(
                    source_file=selected_files[0].resolve().as_posix(), sheet_name="Test",
                    source_row_no=row_no, business_key="PI10087221",
                    source_key=f"detail-{row_no}", row_hash="same-hash", field_values={
                        "erp_order_no": "PI10087221", "erp_internal_product_no": "02PT28329A01",
                    }, source_table="qms_customer_order", field_definitions={},
                ), None

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source"
            source.mkdir()
            (source / "客户订单资料.xls").write_bytes(b"today")
            client = Client()
            config = {"skip_initialization": True, "incremental_directory": str(source)}
            messages = []
            with patch.object(sync, "state_path", return_value=root / "state.json"), \
                    patch.object(sync, "iter_all_source_records", side_effect=fake_records), \
                    patch.object(sync, "QmsDatabaseClient", return_value=client):
                sync.save_state({sync.SOURCE_IDENTITY_VERSION_STATE_KEY: "5"})
                sync.synchronize_database_directory(config, False, messages.append)
                self.assertEqual(["detail-2", "detail-3"], client.written_keys)
                state = sync.load_state()
                self.assertEqual("same-hash", state["raw:detail-2"])
                self.assertEqual("same-hash", state["raw:detail-3"])

    def test_file_replaced_while_reading_does_not_advance_checkpoint(self) -> None:
        class FakeConnection:
            rollbacks = 0

            def cursor(self):
                return self

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            def commit(self):
                raise AssertionError("changed source must not commit")

            def rollback(self):
                self.rollbacks += 1

            def close(self):
                pass

        class FakeClient:
            def __init__(self, connection):
                self.connection = connection

            def _existing_source_tables(self, _cursor):
                return set()

            def _connect(self):
                return self.connection

            def _upsert_source_records(self, *_args):
                pass

        def fake_records(_config, _callback, stats, *, source_directory, selected_files):
            path = selected_files[0]
            stats.files_scanned += 1
            stats.raw_records += 1
            path.write_bytes(b"later")
            yield sync.RawSourceRecord(
                source_file=path.name, sheet_name="Sheet1", source_row_no=2,
                business_key="1", source_key="raw1", row_hash="hash1",
                field_values={}, source_table="qms_material", field_definitions={},
            ), None

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source"
            source.mkdir()
            (source / "today.xls").write_bytes(b"first")
            config = {"skip_initialization": True, "incremental_directory": str(source)}
            connection = FakeConnection()
            with patch.object(sync, "state_path", return_value=root / "state.json"), \
                    patch.object(sync, "iter_all_source_records", side_effect=fake_records), \
                    patch.object(sync, "QmsDatabaseClient", return_value=FakeClient(connection)):
                sync.save_state({sync.SOURCE_IDENTITY_VERSION_STATE_KEY: "5"})
                with self.assertRaisesRegex(sync.SyncError, "覆盖"):
                    sync.synchronize_database_directory(config, False, lambda _: None)
                self.assertEqual({sync.SOURCE_IDENTITY_VERSION_STATE_KEY: "5"}, sync.load_state())
            self.assertEqual(1, connection.rollbacks)

    def test_overwrite_deactivates_rows_removed_from_the_same_file(self) -> None:
        class FakeConnection:
            def cursor(self):
                return self

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            def commit(self):
                pass

            def close(self):
                pass

        class FakeClient:
            def __init__(self):
                self.removed = []

            def _existing_source_tables(self, _cursor):
                return set()

            def _connect(self):
                return FakeConnection()

            def _upsert_source_records(self, *_args):
                pass

            def _upsert_work_order(self, *_args):
                return "work-order-id"

            def _upsert_batch(self, *_args):
                pass

            def _deactivate_removed_file_records(self, _cursor, removed):
                self.removed.append(removed)
                return sum(map(len, removed["raw"].values())) + len(removed["work_orders"]) + len(removed["batches"])

        def fake_records(_config, _callback, stats, *, source_directory, selected_files):
            path = selected_files[0]
            stats.files_scanned += 1
            if path.read_bytes() == b"empty":
                return
            for number in (1, 2) if path.read_bytes() == b"first" else (1,):
                stats.raw_records += 1
                stats.work_order_records += 1
                yield sync.RawSourceRecord(
                    source_file=path.name, sheet_name="Sheet1", source_row_no=number + 1,
                    business_key=str(number), source_key=f"raw{number}", row_hash=f"hash{number}",
                    field_values={}, source_table="qms_purchase_order", field_definitions={},
                ), {"sourceKey": f"ERP-EXCEL:{number}", "workOrderNo": str(number), "batchNo": f"B{number}"}

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source"
            source.mkdir()
            workbook = source / "采购订单资料.xls"
            workbook.write_bytes(b"first")
            config = {"skip_initialization": True, "incremental_directory": str(source)}
            client = FakeClient()
            with patch.object(sync, "state_path", return_value=root / "state.json"), \
                    patch.object(sync, "iter_all_source_records", side_effect=fake_records), \
                    patch.object(sync, "QmsDatabaseClient", return_value=client):
                sync.save_state({sync.SOURCE_IDENTITY_VERSION_STATE_KEY: "5"})
                sync.synchronize_database_directory(config, False, lambda _: None)
                file_key = sync.source_file_state_key("incremental", source, workbook)
                original_membership = sync.load_state()[sync.file_membership_state_key(file_key)]
                previous_deactivations = len(client.removed)
                workbook.write_bytes(b"empty")
                sync.synchronize_database_directory(config, False, lambda _: None)
                self.assertEqual(previous_deactivations, len(client.removed))
                self.assertEqual(
                    original_membership,
                    sync.load_state()[sync.file_membership_state_key(file_key)],
                )
                self.assertEqual(0, sync.synchronize_database_directory(config, False, lambda _: None).source_files)
                workbook.write_bytes(b"later")
                result = sync.synchronize_database_directory(config, False, lambda _: None)
            self.assertEqual(3, result.deactivated)
            self.assertEqual({
                "raw": {"qms_purchase_order": ["raw2"]},
                "work_orders": ["ERP-EXCEL:2"],
                "batches": ["B2"],
            }, client.removed[-1])

    def test_disconnect_preserves_first_file_and_retries_only_second(self) -> None:
        class FakeConnection:
            def __init__(self):
                self.commits = 0
                self.rollbacks = 0

            def cursor(self):
                return self

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            def commit(self):
                self.commits += 1

            def rollback(self):
                self.rollbacks += 1

            def close(self):
                pass

        class FakeClient:
            def __init__(self, _config):
                self.connections = []
                self.fail_second = True

            def _existing_source_tables(self, _cursor):
                return set()

            def _connect(self):
                connection = FakeConnection()
                self.connections.append(connection)
                return connection

            def _upsert_source_records(self, _cursor, _table, records):
                if self.fail_second and any(record.source_file == "second.xls" for record in records):
                    raise pymysql.OperationalError(2013, "connection lost")

            def _deactivate_removed_file_records(self, _cursor, _removed):
                return 0

        def fake_records(_config, _callback, stats, *, source_directory, selected_files):
            for path in selected_files:
                stats.files_scanned += 1
                stats.raw_records += 1
                name = path.name
                yield sync.RawSourceRecord(
                    source_file=name, sheet_name="Sheet1", source_row_no=2,
                    business_key=name, source_key=name, row_hash=name,
                    field_values={}, source_table="qms_material",
                    field_definitions={},
                ), None

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source"
            source.mkdir()
            (source / "first.xls").write_bytes(b"first")
            (source / "second.xls").write_bytes(b"second")
            config = {"skip_initialization": True, "incremental_directory": str(source)}
            first_client = FakeClient(config)
            with patch.object(sync, "state_path", return_value=root / "state.json"), \
                    patch.object(sync, "iter_all_source_records", side_effect=fake_records), \
                    patch.object(sync, "QmsDatabaseClient", return_value=first_client):
                sync.save_state({sync.SOURCE_IDENTITY_VERSION_STATE_KEY: "5"})
                with self.assertRaises(sync.SyncError):
                    sync.synchronize_database_directory(config, False, lambda _: None)
                first_state = sync.load_state()
                first_key = sync.source_file_state_key("incremental", source, source / "first.xls")
                second_key = sync.source_file_state_key("incremental", source, source / "second.xls")
                self.assertIn(first_key, first_state)
                self.assertNotIn(second_key, first_state)
                self.assertEqual(1, first_client.connections[0].commits)
                self.assertEqual(1, first_client.connections[1].rollbacks)

                retry_client = FakeClient(config)
                retry_client.fail_second = False
                with patch.object(sync, "QmsDatabaseClient", return_value=retry_client):
                    result = sync.synchronize_database_directory(config, False, lambda _: None)
                self.assertEqual(1, result.source_files)
                self.assertEqual(1, len(retry_client.connections))
                self.assertIn(second_key, sync.load_state())


if __name__ == "__main__":
    unittest.main()
