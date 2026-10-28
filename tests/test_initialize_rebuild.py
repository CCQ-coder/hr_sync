"""Safety checks for the explicit full rebuild."""

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, Mock, patch

from openpyxl import Workbook

import qms_erp_sync as sync
from pymysql.cursors import RE_INSERT_VALUES


class SourceIdentityTests(unittest.TestCase):
    def test_material_issue_keeps_receipt_details_distinct_across_exports(self):
        columns = (
            ("出库单明细ID", "erp_material_issue_detail_id"),
            ("入库单明细ID编号", "erp_receipt_detail_id_no"),
            ("出库重量", "erp_issue_weight"),
        )

        def record(file_name, row_no, receipt_detail, weight):
            definitions = tuple(sync.ErpFieldDefinition(
                "qms_material_issue", file_name, "Test", header, header, column,
                column, header, index,
            ) for index, (header, column) in enumerate(columns, 1))
            raw_data = dict(zip((header for header, _ in columns),
                                ("ISSUE-1", receipt_detail, weight)))
            return next(iter(sync._source_record_pair(
                file_name, "Test", row_no, raw_data, raw_data, 0, False,
                sync.SourceScanStats(), "qms_material_issue", definitions,
                file_name,
            )))[0]

        first = record("baseline/原料出库单资料.xls", 2, "RECEIPT-1", 10)
        second = record("baseline/原料出库单资料.xls", 3, "RECEIPT-2", 10)
        changed = record("daily/原料出库单资料.xls", 25, "RECEIPT-1", 20)
        self.assertNotEqual(first.source_key, second.source_key)
        self.assertEqual(first.source_key, changed.source_key)
        self.assertNotEqual(first.row_hash, changed.row_hash)

    def test_material_issue_without_receipt_detail_keeps_existing_fallback(self):
        first = {"erp_material_issue_detail_id": "ISSUE-1", "erp_issue_weight": "10"}
        changed = {**first, "erp_issue_weight": "20"}
        self.assertEqual(
            sync.erp_row_identity("qms_material_issue", first),
            sync.erp_row_identity("qms_material_issue", changed),
        )

    def test_finished_goods_receipt_uses_erp_receipt_id_before_shared_receipt_number(self):
        first = {
            "erp_receipt_id_m": "DG01646538", "erp_receipt_no": "DF01173552",
            "erp_internal_product_no": "02PH02148A01", "erp_current_receipt_quantity": "100",
        }
        other_detail = {**first, "erp_receipt_id_m": "DG01646533"}
        updated = {**first, "erp_current_receipt_quantity": "200"}
        self.assertNotEqual(
            sync.erp_row_identity("qms_finished_goods_receipt", first),
            sync.erp_row_identity("qms_finished_goods_receipt", other_detail),
        )
        self.assertEqual(
            sync.erp_row_identity("qms_finished_goods_receipt", first),
            sync.erp_row_identity("qms_finished_goods_receipt", updated),
        )

    def test_raw_upsert_is_compatible_with_pymysql_multirow_insert(self):
        client = sync.QmsDatabaseClient(sync.DEFAULT_CONFIG)
        cursor = Mock()
        records = [sync.RawSourceRecord(
            "a.xls", "Sheet1", number, str(number), f"ERP5:{number}", "hash", {},
            "qms_material", {},
        ) for number in (2, 3)]
        client._upsert_source_records(cursor, "qms_material", records, ensure_schema=False)
        statement, rows = cursor.executemany.call_args.args
        self.assertEqual(2, len(rows))
        self.assertIsNotNone(RE_INSERT_VALUES.match(statement))

    def test_existing_legacy_database_keys_require_rebuild(self):
        class Cursor:
            def execute(self, statement, *_args):
                self.statement = statement

            def fetchone(self):
                return (1,) if "NOT LIKE 'ERP5:%'" in self.statement else None

            def executemany(self, *_args):
                pass

        client = sync.QmsDatabaseClient({
            "database_host": "localhost", "database_port": "3306",
            "database_name": "qms", "database_username": "qms", "tenant_code": "t",
        })
        record = sync.RawSourceRecord(
            "a.xls", "Sheet1", 2, "A", "ERP5:new", "hash", {},
            "qms_material", {},
        )
        with patch.object(client, "_ensure_source_table_schema"):
            with self.assertRaisesRegex(sync.SyncError, "旧版来源键"):
                client._upsert_source_records(Cursor(), "qms_material", [record])

    def test_business_number_updates_across_header_aliases_and_keeps_detail_lines(self):
        def record(file_name, headers, values):
            definitions = tuple(sync.ErpFieldDefinition(
                "qms_material_receipt", file_name, "Sheet1", header, header, column,
                column, header, index,
            ) for index, (header, column) in enumerate(zip(
                headers, ("erp_receipt_no", "erp_material_no", "erp_quantity"),
            ), 1))
            raw_data = dict(zip(headers, values))
            return next(iter(sync._source_record_pair(
                file_name, "Sheet1", 2, raw_data, raw_data, 0, False,
                sync.SourceScanStats(), "qms_material_receipt", definitions, file_name,
            )))[0]

        first = record("baseline.xls", ("入库单号", "原料编号", "数量"), ("R1", "M1", 5))
        changed = record("daily.xls", ("ylr_bh", "yuanl_bh", "shul_mq"), ("R1", "M1", 9))
        other_line = record("daily.xls", ("ylr_bh", "yuanl_bh", "shul_mq"), ("R1", "M2", 9))
        self.assertEqual(first.source_key, changed.source_key)
        self.assertNotEqual(first.row_hash, changed.row_hash)
        self.assertNotEqual(first.source_key, other_line.source_key)

    def test_purchase_order_repeated_material_keeps_each_occurrence(self):
        headers = ["采购单号", "原料编号", "数量", "单价", "交货日期"]
        columns = [
            "erp_purchase_order_no", "erp_material_no", "erp_quantity",
            "erp_unit_price", "erp_delivery_date",
        ]
        registry = sync.ErpFieldRegistry(
            sync.ErpFieldDefinition(
                "qms_purchase_order", "采购订单资料.xlsx", "Test", header, header,
                column, sync.legacy_database_column_name(header),
                sync.erp_column_comment(header), index,
            ) for index, (header, column) in enumerate(zip(headers, columns), 1)
        )
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            exports = []
            for day, rows in (
                ("baseline", ((10, 2, "2026-10-26"), (20, 3, "2026-10-27"))),
                ("daily", ((11, 2, "2026-10-28"), (21, 4, "2026-10-29"))),
            ):
                folder = root / day
                folder.mkdir()
                path = folder / "采购订单资料.xlsx"
                workbook = Workbook()
                workbook.active.title = "Test"
                workbook.active.append(headers)
                for quantity, price, date in rows:
                    workbook.active.append(["KS2504059", "GB10002353", quantity, price, date])
                workbook.save(path)
                workbook.close()
                with patch.object(sync, "load_erp_field_registry", return_value=registry):
                    exports.append([record for record, _ in sync.iter_all_source_records(
                        {}, lambda _: None, sync.SourceScanStats(),
                        source_directory=root, selected_files=(path,),
                    )])

        baseline, daily = exports
        self.assertEqual(2, len({row.source_key for row in baseline}))
        self.assertEqual([row.source_key for row in baseline], [row.source_key for row in daily])
        self.assertNotEqual(baseline[0].row_hash, daily[0].row_hash)

    def test_customer_order_occurrences_keep_identical_rows_and_match_later_exports(self):
        class Registry:
            def canonical_headers_for(self, _source_file, headers):
                return headers

            def resolve_sheet(self, table, source_file, sheet, headers):
                columns = (
                    "erp_order_no", "erp_internal_product_no",
                    "erp_delivery_date", "erp_order_placed_quantity",
                )
                return [sync.ErpFieldDefinition(
                    table, source_file, sheet, header, header, column,
                    column, header, index,
                ) for index, (header, column) in enumerate(zip(headers, columns), 1)]

        with tempfile.TemporaryDirectory() as temporary:
            exports = []
            for day, quantities in (("baseline", (4000, 7025)), ("daily", (5000, 8000))):
                folder = Path(temporary) / day
                folder.mkdir()
                path = folder / "客户订单资料.xlsx"
                workbook = Workbook()
                first = workbook.active
                first.title = "前两行"
                first.append(["订单编号", "厂内编号", "交期", "下单数量"])
                first.append(["PI1", "L1", "2025-01-17", quantities[0]])
                first.append(["PI1", "L1", "2025-03-30", quantities[1]])
                second = workbook.create_sheet("另一工作表")
                second.append(["订单编号", "厂内编号", "交期", "下单数量"])
                second.append(["PI1", "L1", "2025-03-30", quantities[1]])
                workbook.save(path)
                with patch.object(sync, "load_erp_field_registry", return_value=Registry()):
                    exports.append([record for record, _ in sync.iter_all_source_records(
                        {}, lambda _: None, sync.SourceScanStats(),
                        source_directory=folder, selected_files=(path,),
                    )])

        baseline, daily = exports
        self.assertEqual(3, len({row.source_key for row in baseline}))
        self.assertEqual([row.source_key for row in baseline], [row.source_key for row in daily])
        self.assertNotEqual(baseline[0].row_hash, daily[0].row_hash)
        self.assertTrue(all(row.source_key.startswith("ERP5:") for row in baseline))


class InitializationTests(unittest.TestCase):
    def test_selected_excel_rebuild_scans_only_that_file(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            baseline = root / "baseline"
            daily = root / "daily"
            baseline.mkdir()
            daily.mkdir()
            selected = baseline / "原料资料明细.xls"
            selected.write_bytes(b"selected")
            (baseline / "报废.xls").write_bytes(b"other baseline")
            (daily / "报废.xls").write_bytes(b"incremental")
            config = {
                "initialization_directory": str(selected),
                "incremental_directory": str(daily),
            }

            groups, _ = sync.build_source_scan_plan(config, {}, force_all=True)

            self.assertEqual(1, len(groups))
            self.assertEqual(baseline, groups[0].directory)
            self.assertEqual((selected,), groups[0].files)

    def test_snapshot_rejects_export_replaced_during_copy(self):
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "export.xls"
            staged = Path(temporary) / "staged.xls"
            source.write_bytes(b"before")
            real_copy = sync.shutil.copyfile

            def overwrite_after_copy(original, target):
                real_copy(original, target)
                source.write_bytes(b"after")

            with patch.object(sync.shutil, "copyfile", side_effect=overwrite_after_copy):
                with self.assertRaisesRegex(sync.SyncError, "复制期间发生覆盖.*数据库尚未清空"):
                    sync.snapshot_source_file(source, staged)

    def test_xls_reader_uses_staged_file(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "原料入库单资料.XLS"
            staged = root / "staged.XLS"
            source.write_bytes(b"source")
            staged.write_bytes(b"snapshot")
            with patch.object(sync, "load_erp_field_registry", return_value=Mock()), \
                    patch.object(sync, "open_erp_xls", side_effect=ValueError("stop")) as opener:
                with self.assertRaisesRegex(sync.SyncError, "无法打开 Excel"):
                    list(sync.iter_all_source_records(
                        {}, lambda _: None, sync.SourceScanStats(),
                        source_directory=root, selected_files=(source,), read_path=staged,
                    ))
            opener.assert_called_once_with(staged)

    def test_rebuild_imports_fixed_copy_when_export_changes_after_clear(self):
        class Connection:
            def cursor(self):
                return self

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            def execute(self, _statement):
                pass

            def commit(self):
                pass

            def rollback(self):
                pass

            def close(self):
                pass

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source_dir = root / "baseline"
            source_dir.mkdir()
            source = source_dir / "其他资料.xlsx"

            def write_export(value):
                workbook = Workbook()
                workbook.active.title = "Test"
                workbook.active.append(["编号"])
                workbook.active.append([value])
                workbook.save(source)
                workbook.close()

            write_export("A")
            first_checkpoint = sync.source_file_checkpoint(source)
            field = sync.ErpFieldDefinition(
                "qms_erp_unclassified", source.name, "Test", "编号", "编号",
                "erp_no", "erp_no", "Excel列名：编号", 1,
            )
            registry = Mock()
            registry.canonical_headers_for.return_value = ["编号"]
            registry.resolve_sheet.return_value = (field,)

            class Client:
                rows = []
                clear_calls = 0

                def _connect(self):
                    return Connection()

                def _ensure_source_table_schema(self, *_args):
                    pass

                def _clear_target_tables(self, _cursor, _callback):
                    self.clear_calls += 1
                    write_export("B")

                def _upsert_source_records(self, _cursor, _table, records, **_kwargs):
                    self.rows.extend(records)

            client = Client()
            opened = []
            actual_load_workbook = sync.load_workbook

            def track_workbook(path, **kwargs):
                opened.append(Path(path))
                return actual_load_workbook(path, **kwargs)

            with patch.object(sync, "state_path", return_value=root / "state.json"), \
                    patch.object(sync, "QmsDatabaseClient", return_value=client), \
                    patch.object(sync, "load_erp_field_registry", return_value=registry), \
                    patch.object(sync, "load_workbook", side_effect=track_workbook):
                sync.synchronize_database_directory(
                    {"initialization_directory": str(source_dir)}, False,
                    lambda _: None, initialize=True,
                )
                saved = sync.load_state()

            file_key = sync.source_file_state_key("initialization", source_dir, source)
            self.assertEqual(1, client.clear_calls)
            self.assertEqual("A", client.rows[0].field_values["erp_no"])
            self.assertEqual(source.resolve().as_posix(), client.rows[0].source_file)
            self.assertEqual(first_checkpoint, saved[file_key])
            self.assertNotEqual(first_checkpoint, sync.source_file_checkpoint(source))
            self.assertEqual(2, len(opened))
            self.assertEqual(opened[0], opened[1])
            self.assertNotEqual(source, opened[0])
            self.assertFalse(opened[0].exists())

    def test_header_preflight_collects_fields_without_converting_data_rows(self):
        field = sync.ErpFieldDefinition(
            "qms_purchase_order", "采购订单资料.xls", "Test", "采购单号", "采购单号",
            "erp_purchase_order_no", sync.legacy_database_column_name("采购单号"),
            sync.erp_column_comment("采购单号"), 1,
        )
        registry = sync.ErpFieldRegistry([field])
        sheet = Mock(nrows=2)
        sheet.row_values.side_effect = lambda row: ["采购单号"] if row == 0 else self.fail("read data row")
        workbook = Mock(encoding="utf-8", biff_version=80, datemode=0)
        workbook.sheet_names.return_value = ["Test"]
        workbook.sheet_by_name.return_value = sheet
        found = []
        stats = sync.SourceScanStats()
        root = Path("/tmp/erp-preflight")
        with patch.object(sync, "open_erp_xls", return_value=workbook), \
                patch.object(sync, "load_erp_field_registry", return_value=registry):
            rows = list(sync.iter_all_source_records(
                {}, lambda _: None, stats, source_directory=root,
                selected_files=[root / "采购订单资料.xls"], headers_only=True,
                on_sheet=lambda table, fields: found.append((table, fields)),
            ))
        self.assertEqual([], rows)
        self.assertEqual([("qms_purchase_order", {field.database_column: field})], found)
        self.assertEqual((1, 1, 0), (stats.files_scanned, stats.sheets_scanned, stats.raw_records))

    def test_header_preflight_skips_xlsx_data_rows(self):
        field = sync.ErpFieldDefinition(
            "qms_purchase_order", "采购订单资料.xls", "Test", "采购单号", "采购单号",
            "erp_purchase_order_no", sync.legacy_database_column_name("采购单号"),
            sync.erp_column_comment("采购单号"), 1,
        )
        registry = sync.ErpFieldRegistry([field])
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            workbook = Workbook()
            workbook.active.title = "Test"
            workbook.active.append(["采购单号"])
            workbook.active.append(["PO-1"])
            path = root / "采购订单资料.xlsx"
            workbook.save(path)
            workbook.close()
            found = []
            stats = sync.SourceScanStats()
            with patch.object(sync, "load_erp_field_registry", return_value=registry):
                rows = list(sync.iter_all_source_records(
                    {}, lambda _: None, stats, source_directory=root,
                    selected_files=[path], headers_only=True,
                    on_sheet=lambda table, fields: found.append((table, fields)),
                ))
        self.assertEqual([], rows)
        self.assertEqual([("qms_purchase_order", {field.database_column: field})], found)
        self.assertEqual(0, stats.raw_records)

    def test_clear_rejects_nontransactional_tables_and_external_dependents(self):
        class Cursor:
            def __init__(self, answers):
                self.answers = iter(answers)
                self.statements = []

            def execute(self, statement, *_args):
                self.statements.append(statement)

            def fetchall(self):
                return next(self.answers)

        client = object.__new__(sync.QmsDatabaseClient)
        nontransactional = Cursor([[("qms_material", "MyISAM")]])
        with self.assertRaisesRegex(sync.SyncError, "InnoDB"):
            client._clear_target_tables(nontransactional, lambda _: None)
        self.assertFalse(any("DELETE FROM" in sql for sql in nontransactional.statements))

        external_reference = Cursor([[("qms_material", "InnoDB")], [("other_table",)]])
        with self.assertRaisesRegex(sync.SyncError, "范围外"):
            client._clear_target_tables(external_reference, lambda _: None)
        self.assertFalse(any("DELETE FROM" in sql for sql in external_reference.statements))

    def test_clear_rejects_tables_without_a_single_id_primary_key(self):
        class Cursor:
            def __init__(self, primary_columns):
                self.answers = iter(([("qms_material", "InnoDB")], [], primary_columns))
                self.statements = []

            def execute(self, statement, *_args):
                self.statements.append(statement)

            def fetchall(self):
                return next(self.answers)

        client = object.__new__(sync.QmsDatabaseClient)
        for primary_columns in ([], [("qms_material", "other")],
                                [("qms_material", "id"), ("qms_material", "other")]):
            with self.subTest(primary_columns=primary_columns):
                cursor = Cursor(primary_columns)
                with self.assertRaisesRegex(sync.SyncError, "id.*主键"):
                    client._clear_target_tables(cursor, lambda _: None)
                self.assertFalse(any("DELETE FROM" in sql for sql in cursor.statements))

    def test_missing_local_state_requires_rebuild(self):
        with patch.object(sync, "load_state", return_value={}):
            with self.assertRaisesRegex(sync.SyncError, "清空并重新初始化"):
                sync.synchronize_database_directory({}, False, lambda _: None)

    def test_rebuild_reads_unchanged_baseline_files(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "原料资料明细.xls"
            source.write_bytes(b"export")
            config = {"initialization_directory": str(root), "skip_initialization": False}
            _, previous = sync.build_source_scan_plan(config, {})
            groups, _ = sync.build_source_scan_plan(config, previous, force_all=True)
            self.assertEqual([source], list(groups[0].files))

    def test_clear_deletes_all_tenants_in_target_tables(self):
        statements = []
        progress = []
        rows = {
            "`qms_production_batch`": [f"{number:04d}" for number in range(5)],
            "`qms_work_order`": [f"{number:04d}" for number in range(12)],
            "`qms_material`": [f"{number:04d}" for number in range(2002)],
        }

        class Cursor:
            def execute(self, statement, parameters=()):
                statements.append((statement, parameters))
                table = statement.split("FROM ")[1].split()[0]
                if statement.startswith("SELECT"):
                    lower = parameters[0] if parameters else None
                    self.selected = [row for row in rows[table] if lower is None or row > lower][:2000]
                    return
                if not parameters:
                    raise AssertionError("unbounded DELETE")
                if "WHERE `id` IN" not in statement or statement.count("%s") != len(parameters):
                    raise AssertionError("DELETE must target the selected primary keys")
                selected = set(parameters)
                remaining = [row for row in rows[table] if row not in selected]
                self.rowcount = len(rows[table]) - len(remaining)
                rows[table] = remaining

            def fetchall(self):
                return [(row,) for row in self.selected]

        client = object.__new__(sync.QmsDatabaseClient)
        with patch.object(client, "_existing_source_tables", return_value={"qms_material"}), \
                patch.object(client, "_assert_safe_clear"):
            client._clear_target_tables(Cursor(), progress.append)
        self.assertTrue(all(not table_rows for table_rows in rows.values()))
        material_calls = [call for call in statements if "`qms_material`" in call[0]]
        self.assertEqual(5, len(material_calls))
        self.assertEqual(
            ("SELECT `id` FROM `qms_material` ORDER BY `id` LIMIT 2000 FOR UPDATE", ()),
            material_calls[0],
        )
        self.assertEqual(("0000", "1999"), (material_calls[1][1][0], material_calls[1][1][-1]))
        self.assertEqual(2000, len(material_calls[1][1]))
        self.assertEqual(
            ("SELECT `id` FROM `qms_material` WHERE `id` > %s ORDER BY `id` LIMIT 2000 FOR UPDATE", ("1999",)),
            material_calls[2],
        )
        self.assertEqual(("2000", "2001"), material_calls[3][1])
        self.assertEqual(
            ("SELECT `id` FROM `qms_material` WHERE `id` > %s ORDER BY `id` LIMIT 2000 FOR UPDATE", ("2001",)),
            material_calls[4],
        )
        self.assertEqual([
            "正在清空目标表 qms_production_batch…",
            "目标表 qms_production_batch 已删除 5 行（待初始化完成后提交）。",
            "正在清空目标表 qms_work_order…",
            "目标表 qms_work_order 已删除 12 行（待初始化完成后提交）。",
            "正在清空目标表 qms_material…",
            "目标表 qms_material 已删除 2,000 行（待初始化完成后提交）…",
            "目标表 qms_material 已删除 2,002 行（待初始化完成后提交）。",
        ], progress)

    def test_selected_partial_export_clears_only_its_own_file_rows(self):
        selected = "/exports/原料出库单资料.xls"
        rows = {
            "a": ("tenant-a", selected),
            "b": ("tenant-a", "/exports/补料资料.xls"),
            "c": ("tenant-b", selected),
            "d": ("tenant-b", "/exports/补料资料.xls"),
        }
        statements = []

        class Cursor:
            def execute(self, statement, parameters=()):
                statements.append((statement, parameters))
                if statement.startswith("SELECT DISTINCT"):
                    self.selected = [(tenant,) for tenant in sorted({tenant for tenant, _ in rows.values()})]
                elif statement.startswith("SELECT"):
                    if "`tenant_code` = %s AND `source_file` = %s" not in statement:
                        raise AssertionError("must use the existing tenant/source_file index")
                    tenant, source = parameters[:2]
                    last_id = parameters[2] if len(parameters) > 2 else None
                    self.selected = [(row_id,) for row_id, values in sorted(rows.items())
                                     if values == (tenant, source) and (last_id is None or row_id > last_id)]
                else:
                    self.rowcount = len(parameters)
                    for row_id in parameters:
                        rows.pop(row_id)

            def fetchall(self):
                return self.selected

        client = object.__new__(sync.QmsDatabaseClient)
        client.tenant = "tenant-a"
        with patch.object(client, "_existing_source_tables", return_value={
            "qms_material_issue", "qms_material",
        }), patch.object(client, "_assert_safe_clear"):
            client._clear_target_tables(
                Cursor(), lambda _: None,
                selected_tables={"qms_material_issue"}, source_file=selected,
            )

        self.assertEqual({
            "b": ("tenant-a", "/exports/补料资料.xls"),
            "d": ("tenant-b", "/exports/补料资料.xls"),
        }, rows)
        self.assertEqual(7, len(statements))
        self.assertTrue(all("`qms_material_issue`" in statement for statement, _ in statements))
        self.assertEqual(("tenant-a", selected), statements[1][1])
        self.assertEqual(("tenant-b", selected), statements[4][1])

        selected_receipt = "/exports/成品入库单资料.xls"
        rows.update({
            "a": ("tenant-a", selected_receipt),
            "c": ("tenant-b", selected_receipt),
        })
        statements.clear()
        with patch.object(client, "_existing_source_tables", return_value={
            "qms_finished_goods_receipt",
        }), patch.object(client, "_assert_safe_clear"):
            client._clear_target_tables(
                Cursor(), lambda _: None,
                selected_tables={"qms_finished_goods_receipt"}, source_file=selected_receipt,
            )
        self.assertEqual({
            "b": ("tenant-a", "/exports/补料资料.xls"),
            "d": ("tenant-b", "/exports/补料资料.xls"),
        }, rows)
        self.assertTrue(all("`qms_finished_goods_receipt`" in statement for statement, _ in statements))

    def test_selected_partial_export_requires_source_file_before_clear(self):
        client = object.__new__(sync.QmsDatabaseClient)
        cursor = Mock()
        with patch.object(client, "_assert_safe_clear"), \
                patch.object(client, "_existing_source_tables", return_value=set()):
            with self.assertRaisesRegex(sync.SyncError, "来源文件"):
                client._clear_target_tables(
                    cursor, lambda _: None, selected_tables={"qms_material_issue"},
                )
        cursor.execute.assert_not_called()


    def test_selected_complete_export_clears_whole_target_table(self):
        statements = []

        class Cursor:
            def execute(self, statement, parameters=()):
                statements.append((statement, parameters))
                self.rowcount = 2

            def fetchall(self):
                if len(statements) == 1:
                    return [("tenant-a-row",), ("tenant-b-row",)]
                return []

        client = object.__new__(sync.QmsDatabaseClient)
        with patch.object(client, "_existing_source_tables", return_value={"qms_material"}), \
                patch.object(client, "_assert_safe_clear"):
            client._clear_target_tables(
                Cursor(), lambda _: None,
                selected_tables={"qms_material"}, source_file="/exports/原料资料明细.xls",
            )

        self.assertEqual(3, len(statements))
        self.assertEqual((), statements[0][1])
        self.assertNotIn("source_file", statements[0][0])
        self.assertEqual(("tenant-a-row", "tenant-b-row"), statements[1][1])


    def test_selected_excel_rebuild_preserves_other_state(self):
        cleared = []
        saved = []
        visited = []

        class Connection:
            def cursor(self):
                return self

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            def execute(self, _statement):
                pass

            def commit(self):
                pass

            def rollback(self):
                pass

            def close(self):
                pass

        class Client:
            def __init__(self, _config):
                pass

            def _connect(self):
                return Connection()

            def _ensure_source_table_schema(self, *_args):
                pass

            def _clear_target_tables(self, _cursor, _callback, **kwargs):
                cleared.append(kwargs)

            def _upsert_source_records(self, *_args, **_kwargs):
                pass

            def _delete_missing_snapshot_rows(self, *_args):
                return 0, set()

        def records(_config, _callback, stats, *, source_directory, selected_files,
                    headers_only=False, on_sheet=None, read_path=None):
            visited.append(selected_files[0])
            stats.files_scanned += 1
            if on_sheet is not None:
                on_sheet("qms_material", {})
            if headers_only:
                return
            stats.raw_records += 1
            yield sync.RawSourceRecord(
                source_file=selected_files[0].resolve().as_posix(), sheet_name="Test", source_row_no=2,
                business_key="MAT-1", source_key="ERP5:MAT-1", row_hash="hash",
                field_values={}, source_table="qms_material", field_definitions={},
            ), None

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "原料资料明细.xls"
            source.write_bytes(b"material")
            daily = root / "daily"
            daily.mkdir()
            (daily / "报废.xls").write_bytes(b"scrap")
            config = {
                "initialization_directory": str(source), "incremental_directory": str(daily),
            }
            file_key = sync.source_file_state_key("initialization", root, source)
            old_state = {
                "unrelated-checkpoint": "keep",
                sync.SOURCE_IDENTITY_VERSION_STATE_KEY: "4",
                sync.COMPLETE_SNAPSHOT_VERSION_STATE_KEY: "pending",
                sync.file_membership_state_key(file_key): json.dumps({
                    "raw": {"qms_material": ["ERP5:OLD"]},
                    "work_orders": ["ERP-EXCEL:OLD"], "batches": [],
                }),
                "raw:ERP5:OLD": "old-hash",
                "work-order:ERP-EXCEL:OLD": "old-hash",
            }
            with patch.object(sync, "load_state", return_value=old_state), \
                    patch.object(sync, "QmsDatabaseClient", Client), \
                    patch.object(sync, "iter_all_source_records", side_effect=records), \
                    patch.object(sync, "save_state", side_effect=lambda state: saved.append(dict(state))), \
                    patch.object(sync, "clear_file_progress"), patch.object(sync, "log_line"):
                sync.synchronize_database_directory(config, False, lambda _: None, initialize=True)

        self.assertEqual([source, source], visited)
        self.assertEqual({"qms_material"}, cleared[0]["selected_tables"])
        self.assertEqual(source.resolve().as_posix(), cleared[0]["source_file"])
        self.assertEqual("keep", saved[-1]["unrelated-checkpoint"])
        self.assertEqual("4", saved[-1][sync.SOURCE_IDENTITY_VERSION_STATE_KEY])
        self.assertEqual("pending", saved[-1][sync.COMPLETE_SNAPSHOT_VERSION_STATE_KEY])
        self.assertEqual("1", saved[-1][sync.BASELINE_REPLAY_STATE_KEY])
        self.assertNotIn("raw:ERP5:OLD", saved[-1])
        self.assertNotIn("work-order:ERP-EXCEL:OLD", saved[-1])
        self.assertEqual("hash", saved[-1]["raw:ERP5:MAT-1"])

    def test_selected_production_instruction_clears_work_order_and_batch(self):
        client = MagicMock()
        client._upsert_work_order.return_value = "work-order-id"

        def records(_config, _callback, stats, *, source_directory, selected_files,
                    headers_only=False, on_sheet=None, read_path=None):
            if on_sheet is not None:
                on_sheet("qms_work_order", {})
            if headers_only:
                return
            stats.raw_records += 1
            yield sync.RawSourceRecord(
                source_file=selected_files[0].resolve().as_posix(), sheet_name="Test", source_row_no=2,
                business_key="WO-1", source_key="ERP5:WO-1", row_hash="hash",
                field_values={}, source_table=None, field_definitions={},
            ), {"sourceKey": "ERP-EXCEL:WO-1", "workOrderNo": "WO-1", "batchNo": "B-1"}

        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "生产指示资料.xls"
            source.write_bytes(b"production")
            with patch.object(sync, "QmsDatabaseClient", return_value=client), \
                    patch.object(sync, "iter_all_source_records", side_effect=records), \
                    patch.object(sync, "save_state"), patch.object(sync, "clear_file_progress"), \
                    patch.object(sync, "log_line"):
                sync.synchronize_database_directory(
                    {"initialization_directory": str(source)}, False, lambda _: None,
                    initialize=True,
                )

        self.assertEqual(
            {"qms_work_order", "qms_production_batch"},
            client._clear_target_tables.call_args.kwargs["selected_tables"],
        )


    def test_empty_selected_excel_never_reaches_database_clear(self):
        class Client:
            def __init__(self, _config):
                pass

            def _connect(self):
                raise AssertionError("empty Excel must stop before database access")

        def no_records(_config, _callback, _stats, *, source_directory,
                       selected_files, headers_only=False, on_sheet=None, read_path=None):
            if headers_only:
                on_sheet("qms_material", {})
            return iter(())

        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "原料资料明细.xls"
            source.write_bytes(b"empty")
            with patch.object(sync, "QmsDatabaseClient", Client), \
                    patch.object(sync, "iter_all_source_records", side_effect=no_records), \
                    patch.object(sync, "save_state") as save_state:
                with self.assertRaisesRegex(sync.SyncError, "没有有效数据"):
                    sync.synchronize_database_directory(
                        {"initialization_directory": str(source)}, False,
                        lambda _: None, initialize=True,
                    )
            save_state.assert_not_called()




    def test_old_location_key_state_requires_rebuild(self):
        with patch.object(sync, "load_state", return_value={
            sync.SOURCE_IDENTITY_VERSION_STATE_KEY: "2", "raw:old-key": "old-hash",
        }):
            with self.assertRaisesRegex(sync.SyncError, "清空并重新初始化"):
                sync.synchronize_database_directory({}, False, lambda _: None)

    def test_failed_rebuild_rolls_back_clear_and_does_not_save_progress(self):
        events = []

        class Connection:
            def cursor(self):
                return self

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            def commit(self):
                events.append("commit")

            def execute(self, statement):
                events.append(statement)

            def rollback(self):
                events.append("rollback")

            def close(self):
                pass

        class Client:
            def __init__(self, _config):
                pass

            def _connect(self):
                return Connection()

            def _ensure_source_table_schema(self, *_args):
                events.append("schema")

            def _clear_target_tables(self, _cursor, _callback):
                events.append("clear")

            def _existing_source_tables(self, _cursor):
                return {"qms_material"}

            def _upsert_source_records(self, *_args, **_kwargs):
                raise sync.SyncError("write failed")

        def records(_config, _callback, stats, *, source_directory, selected_files,
                    headers_only=False, on_sheet=None, read_path=None):
            stats.files_scanned += 1
            if headers_only:
                on_sheet("qms_material", {})
                return
            stats.raw_records += 1
            yield sync.RawSourceRecord(
                source_file=selected_files[0].name, sheet_name="Sheet1", source_row_no=2,
                business_key="A", source_key="key", row_hash="hash", field_values={},
                source_table="qms_material", field_definitions={},
            ), None

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "baseline"
            source.mkdir()
            (source / "原料资料明细.xls").write_bytes(b"export")
            config = {"initialization_directory": str(source), "skip_initialization": False}
            with patch.object(sync, "state_path", return_value=root / "state.json"), \
                    patch.object(sync, "QmsDatabaseClient", Client), \
                    patch.object(sync, "iter_all_source_records", side_effect=records), \
                    patch.object(sync, "save_state") as save_state, \
                    patch.object(sync, "save_file_progress") as save_progress:
                with self.assertRaisesRegex(sync.SyncError, "write failed"):
                    sync.synchronize_database_directory(config, False, lambda _: None, initialize=True)
            self.assertEqual([
                "schema", "commit", "SET TRANSACTION ISOLATION LEVEL REPEATABLE READ",
                "clear", "rollback",
            ], events)
            save_state.assert_not_called()
            save_progress.assert_not_called()

    def test_rebuild_replaces_older_complete_export_with_latest_in_one_transaction(self):
        events = []

        class Connection:
            def cursor(self):
                return self

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            def commit(self):
                events.append("commit")

            def execute(self, statement):
                events.append(statement)

            def rollback(self):
                events.append("rollback")

            def close(self):
                pass

        class Client:
            def __init__(self, _config):
                pass

            def _connect(self):
                return Connection()

            def _ensure_source_table_schema(self, *_args):
                events.append("schema")

            def _clear_target_tables(self, _cursor, _callback):
                events.append("clear")

            def _upsert_source_records(self, _cursor, _table, rows, **_kwargs):
                events.append(("upsert", rows[0].source_key))

            def _delete_missing_snapshot_rows(self, _cursor, table, keys):
                events.append(("delete", table, set(keys)))
                return 0, set()

        def records(_config, _callback, stats, *, source_directory, selected_files,
                    headers_only=False, on_sheet=None, read_path=None):
            stats.files_scanned += 1
            if headers_only:
                on_sheet("qms_material", {})
                return
            key = (read_path or selected_files[0]).read_text()
            if key == "empty":
                return
            stats.raw_records += 1
            yield sync.RawSourceRecord(
                source_file=selected_files[0].as_posix(), sheet_name="Test", source_row_no=2,
                business_key=key, source_key=f"ERP5:{key}", row_hash=key,
                field_values={}, source_table="qms_material", field_definitions={},
            ), None

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            baseline = root / "baseline"
            for day, value in (("2026-10-22", "old"), ("2026-10-23", "new"),
                               ("2026-10-27", "empty")):
                folder = baseline / day
                folder.mkdir(parents=True)
                (folder / "原料资料明细.xls").write_text(value)
            config = {"initialization_directory": str(baseline)}
            messages = []
            with patch.object(sync, "state_path", return_value=root / "state.json"), \
                    patch.object(sync, "QmsDatabaseClient", Client), \
                    patch.object(sync, "iter_all_source_records", side_effect=records):
                sync.synchronize_database_directory(config, False, messages.append, initialize=True)
                state = sync.load_state()
            self.assertEqual([
                "schema", "commit", "SET TRANSACTION ISOLATION LEVEL REPEATABLE READ", "clear",
                ("upsert", "ERP5:old"),
                ("upsert", "ERP5:new"), ("delete", "qms_material", {"ERP5:new"}),
                "commit",
            ], events)
            self.assertEqual("1", state["complete-snapshot-version"])
            self.assertTrue(any("跳过没有有效数据" in message for message in messages))


if __name__ == "__main__":
    unittest.main()
